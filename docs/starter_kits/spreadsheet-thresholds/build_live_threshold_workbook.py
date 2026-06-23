"""Build the live Week 3 threshold workbook.

The workbook is the classroom object: change thresholds in the Criteria sheet,
watch pass/review/fail colors update, then export Variant_Scores as
variant_scores.csv for Grasshopper.
"""

import csv
from html import escape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "decision_matrix_filled_example.csv"
OUTPUT = ROOT / "live_threshold_workbook.xlsx"
PREVIEW = ROOT / "live_threshold_workbook_preview.html"


FILLS = {
    "header": PatternFill("solid", fgColor="243447"),
    "input": PatternFill("solid", fgColor="FFF2CC"),
    "pass": PatternFill("solid", fgColor="D9EAD3"),
    "review": PatternFill("solid", fgColor="FCE5CD"),
    "fail": PatternFill("solid", fgColor="F4CCCC"),
    "note": PatternFill("solid", fgColor="EAF2F8"),
}


CRITERIA = [
    {
        "criterion": "morning glare",
        "unit": "h/day",
        "direction": "<=",
        "threshold": 2.0,
        "review_margin": 0.15,
        "weight": 0.25,
        "prompt": "How much direct glare is acceptable before the workspace claim weakens?",
    },
    {
        "criterion": "view openness",
        "unit": "fraction",
        "direction": ">=",
        "threshold": 0.55,
        "review_margin": 0.05,
        "weight": 0.20,
        "prompt": "How much visual openness must remain after shading or planting?",
    },
    {
        "criterion": "planting depth",
        "unit": "m",
        "direction": ">=",
        "threshold": 0.80,
        "review_margin": 0.05,
        "weight": 0.15,
        "prompt": "What minimum soil depth makes the greenery claim credible?",
    },
    {
        "criterion": "solar gain proxy",
        "unit": "index",
        "direction": "<=",
        "threshold": 48.0,
        "review_margin": 0.10,
        "weight": 0.25,
        "prompt": "How low must the heat proxy be before the facade strategy is worth developing?",
    },
    {
        "criterion": "added facade depth",
        "unit": "m",
        "direction": "<=",
        "threshold": 0.75,
        "review_margin": 0.10,
        "weight": 0.15,
        "prompt": "How much facade depth can the project tolerate before the solution creates a new problem?",
    },
]

SCORE_RULES = {
    "candidate_min_score": 80,
    "revise_if_fail_count_at_least": 2,
    "review_score_floor": 50,
}


def style_header(row):
    for cell in row:
        cell.fill = FILLS["header"]
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def set_widths(ws, widths):
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def add_table(ws, name, ref):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def build_workbook() -> None:
    with SOURCE.open(newline="") as handle:
        records = list(csv.DictReader(handle))
    for record in records:
        record["value"] = float(record["value"])
    variants = []
    for record in records:
        if record["variant"] not in variants:
            variants.append(record["variant"])

    wb = Workbook()
    demo = wb.active
    demo.title = "Live_Demo"
    criteria = wb.create_sheet("Criteria")
    rules = wb.create_sheet("Score_Rules")
    matrix = wb.create_sheet("Decision_Matrix")
    scores = wb.create_sheet("Variant_Scores")

    build_demo_sheet(demo)
    build_criteria_sheet(criteria)
    build_rules_sheet(rules)
    build_matrix_sheet(matrix, records)
    build_scores_sheet(scores, variants)

    try:
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except AttributeError:
        pass

    wb.save(OUTPUT)
    build_preview_html(records, variants)


def build_demo_sheet(ws) -> None:
    ws["A1"] = "Week 3 live threshold workbook"
    ws["A1"].font = Font(size=18, bold=True, color="243447")
    ws["A3"] = "Use this live:"
    ws["A3"].font = Font(bold=True)
    instructions = [
        "1. Open the Criteria sheet.",
        "2. Change the yellow threshold cells.",
        "3. Watch Decision_Matrix update pass/review/fail colors.",
        "4. Open Score_Rules and change the candidate score or failure limit.",
        "5. Watch Variant_Scores update the status, rank, and chart.",
        "6. Export Variant_Scores as variant_scores.csv for Grasshopper.",
    ]
    for idx, text in enumerate(instructions, start=4):
        ws[f"A{idx}"] = text
    ws["A11"] = "Teaching question"
    ws["A11"].font = Font(bold=True)
    ws["A12"] = "Which threshold is actually a design judgment rather than a fact?"
    ws["A12"].fill = FILLS["note"]
    ws["A12"].alignment = Alignment(wrap_text=True)
    set_widths(ws, {"A": 88})


def build_criteria_sheet(ws) -> None:
    headers = ["criterion", "unit", "direction", "threshold", "review_margin", "weight", "discussion_prompt"]
    ws.append(headers)
    style_header(ws[1])
    for item in CRITERIA:
        ws.append(
            [
                item["criterion"],
                item["unit"],
                item["direction"],
                item["threshold"],
                item["review_margin"],
                item["weight"],
                item["prompt"],
            ]
        )
    for row in range(2, len(CRITERIA) + 2):
        ws[f"D{row}"].fill = FILLS["input"]
        ws[f"E{row}"].fill = FILLS["input"]
        ws[f"F{row}"].fill = FILLS["input"]
        ws[f"E{row}"].number_format = "0%"
        ws[f"F{row}"].number_format = "0%"
        ws[f"G{row}"].alignment = Alignment(wrap_text=True)
    direction_validation = DataValidation(type="list", formula1='"<=,>="')
    ws.add_data_validation(direction_validation)
    direction_validation.add(f"C2:C{len(CRITERIA) + 1}")
    add_table(ws, "CriteriaTable", f"A1:G{len(CRITERIA) + 1}")
    set_widths(ws, {"A": 20, "B": 12, "C": 12, "D": 12, "E": 15, "F": 10, "G": 72})
    ws.freeze_panes = "A2"


def build_rules_sheet(ws) -> None:
    headers = ["rule", "value", "what changes when edited"]
    ws.append(headers)
    style_header(ws[1])
    rows = [
        (
            "candidate_min_score",
            SCORE_RULES["candidate_min_score"],
            "Minimum weighted score for a variant to count as a candidate.",
        ),
        (
            "revise_if_fail_count_at_least",
            SCORE_RULES["revise_if_fail_count_at_least"],
            "A variant with this many failed criteria is forced to revise, even if the score looks high.",
        ),
        (
            "review_score_floor",
            SCORE_RULES["review_score_floor"],
            "Below this score, a non-candidate is treated as revise rather than review.",
        ),
    ]
    for row_idx, row in enumerate(rows, start=2):
        ws.append(row)
        ws[f"B{row_idx}"].fill = FILLS["input"]
        ws[f"C{row_idx}"].alignment = Alignment(wrap_text=True)
    add_table(ws, "ScoreRulesTable", "A1:C4")
    set_widths(ws, {"A": 32, "B": 14, "C": 76})
    ws.freeze_panes = "A2"


def build_matrix_sheet(ws, records) -> None:
    headers = [
        "variant",
        "criterion",
        "value",
        "unit",
        "source",
        "confidence",
        "direction",
        "threshold",
        "review_margin",
        "weight",
        "live_status",
        "weighted_points",
        "source_note",
    ]
    ws.append(headers)
    style_header(ws[1])

    confidence_validation = DataValidation(type="list", formula1='"high,medium,low"')
    ws.add_data_validation(confidence_validation)

    for idx, record in enumerate(records, start=2):
        ws.append(
            [
                record["variant"],
                record["criterion"],
                record["value"],
                record["unit"],
                record["source"],
                record["confidence"],
                f'=VLOOKUP(B{idx},Criteria!$A:$G,3,FALSE)',
                f'=VLOOKUP(B{idx},Criteria!$A:$G,4,FALSE)',
                f'=VLOOKUP(B{idx},Criteria!$A:$G,5,FALSE)',
                f'=VLOOKUP(B{idx},Criteria!$A:$G,6,FALSE)',
                (
                    f'=IF(G{idx}="<=",'
                    f'IF(C{idx}<=H{idx},"pass",IF(C{idx}<=H{idx}*(1+I{idx}),"review","fail")),'
                    f'IF(C{idx}>=H{idx},"pass",IF(C{idx}>=H{idx}*(1-I{idx}),"review","fail")))'
                ),
                f'=IF(K{idx}="pass",J{idx},IF(K{idx}="review",J{idx}*0.5,0))',
                record["decision_note"],
            ]
        )
        confidence_validation.add(ws[f"F{idx}"])
        ws[f"I{idx}"].number_format = "0%"
        ws[f"J{idx}"].number_format = "0%"
        ws[f"L{idx}"].number_format = "0%"
        ws[f"M{idx}"].alignment = Alignment(wrap_text=True)

    last_row = len(records) + 1
    ws.conditional_formatting.add(f"K2:K{last_row}", FormulaRule(formula=['K2="pass"'], fill=FILLS["pass"]))
    ws.conditional_formatting.add(f"K2:K{last_row}", FormulaRule(formula=['K2="review"'], fill=FILLS["review"]))
    ws.conditional_formatting.add(f"K2:K{last_row}", FormulaRule(formula=['K2="fail"'], fill=FILLS["fail"]))
    add_table(ws, "DecisionMatrixTable", f"A1:M{last_row}")
    set_widths(
        ws,
        {
            "A": 18,
            "B": 18,
            "C": 10,
            "D": 12,
            "E": 28,
            "F": 12,
            "G": 10,
            "H": 12,
            "I": 15,
            "J": 10,
            "K": 14,
            "L": 16,
            "M": 48,
        },
    )
    ws.freeze_panes = "A2"


def build_scores_sheet(ws, variants) -> None:
    headers = [
        "variant",
        "pass_count",
        "review_count",
        "fail_count",
        "total_criteria",
        "score_pct",
        "status",
        "rank",
        "decision_note",
        "gh_action",
    ]
    ws.append(headers)
    style_header(ws[1])

    for idx, variant in enumerate(variants, start=2):
        ws.append(
            [
                variant,
                f'=COUNTIFS(Decision_Matrix!$A:$A,A{idx},Decision_Matrix!$K:$K,"pass")',
                f'=COUNTIFS(Decision_Matrix!$A:$A,A{idx},Decision_Matrix!$K:$K,"review")',
                f'=COUNTIFS(Decision_Matrix!$A:$A,A{idx},Decision_Matrix!$K:$K,"fail")',
                f"=SUM(B{idx}:D{idx})",
                (
                    f'=ROUND(SUMIF(Decision_Matrix!$A:$A,A{idx},Decision_Matrix!$L:$L)'
                    f'/SUMIF(Decision_Matrix!$A:$A,A{idx},Decision_Matrix!$J:$J)*100,0)'
                ),
                (
                    f'=IF(D{idx}>=Score_Rules!$B$3,"revise",'
                    f'IF(F{idx}>=Score_Rules!$B$2,"candidate",'
                    f'IF(F{idx}>=Score_Rules!$B$4,"review","revise")))'
                ),
                f"=RANK(F{idx},$F$2:$F${len(variants) + 1},0)",
                f'=IF(G{idx}="candidate","Keep as design-development candidate",IF(G{idx}="review","Revise or defend the borderline criterion","Revise before design development"))',
                f'=IF(G{idx}="candidate","color green and keep",IF(G{idx}="review","color amber and test revision","color red or cull from candidate set"))',
            ]
        )
        ws[f"F{idx}"].number_format = "0"
        ws[f"I{idx}"].alignment = Alignment(wrap_text=True)
        ws[f"J{idx}"].alignment = Alignment(wrap_text=True)

    last_row = len(variants) + 1
    ws.conditional_formatting.add(
        f"F2:F{last_row}",
        CellIsRule(operator="greaterThanOrEqual", formula=[str(SCORE_RULES["candidate_min_score"])], fill=FILLS["pass"]),
    )
    ws.conditional_formatting.add(
        f"F2:F{last_row}",
        CellIsRule(
            operator="between",
            formula=[str(SCORE_RULES["review_score_floor"]), str(SCORE_RULES["candidate_min_score"] - 1)],
            fill=FILLS["review"],
        ),
    )
    ws.conditional_formatting.add(
        f"F2:F{last_row}",
        CellIsRule(operator="lessThan", formula=[str(SCORE_RULES["review_score_floor"])], fill=FILLS["fail"]),
    )
    ws.conditional_formatting.add(f"G2:G{last_row}", FormulaRule(formula=['G2="candidate"'], fill=FILLS["pass"]))
    ws.conditional_formatting.add(f"G2:G{last_row}", FormulaRule(formula=['G2="review"'], fill=FILLS["review"]))
    ws.conditional_formatting.add(f"G2:G{last_row}", FormulaRule(formula=['G2="revise"'], fill=FILLS["fail"]))

    add_table(ws, "VariantScoresTable", f"A1:J{last_row}")
    set_widths(ws, {"A": 18, "B": 12, "C": 14, "D": 12, "E": 14, "F": 12, "G": 14, "H": 8, "I": 42, "J": 42})
    ws.freeze_panes = "A2"

    chart = BarChart()
    chart.title = "Variant score by live threshold"
    chart.y_axis.title = "score percent"
    chart.x_axis.title = "variant"
    data = Reference(ws, min_col=6, min_row=1, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 8
    chart.width = 14
    ws.add_chart(chart, "K2")


def summarize_variant(records, variant):
    criteria_lookup = {item["criterion"]: item for item in CRITERIA}
    pass_count = review_count = fail_count = 0
    weighted = 0.0
    total_weight = 0.0
    for record in records:
        if record["variant"] != variant:
            continue
        criterion = criteria_lookup[record["criterion"]]
        value = float(record["value"])
        threshold = criterion["threshold"]
        margin = criterion["review_margin"]
        weight = criterion["weight"]
        total_weight += weight
        if criterion["direction"] == "<=":
            if value <= threshold:
                status = "pass"
            elif value <= threshold * (1 + margin):
                status = "review"
            else:
                status = "fail"
        else:
            if value >= threshold:
                status = "pass"
            elif value >= threshold * (1 - margin):
                status = "review"
            else:
                status = "fail"
        pass_count += status == "pass"
        review_count += status == "review"
        fail_count += status == "fail"
        weighted += weight if status == "pass" else weight * 0.5 if status == "review" else 0
    score = round(weighted / total_weight * 100)
    if fail_count >= SCORE_RULES["revise_if_fail_count_at_least"]:
        overall = "revise"
    elif score >= SCORE_RULES["candidate_min_score"]:
        overall = "candidate"
    elif score >= SCORE_RULES["review_score_floor"]:
        overall = "review"
    else:
        overall = "revise"
    return {
        "variant": variant,
        "pass": pass_count,
        "review": review_count,
        "fail": fail_count,
        "score": score,
        "status": overall,
    }


def build_preview_html(records, variants) -> None:
    summaries = [summarize_variant(records, variant) for variant in variants]
    summaries = sorted(summaries, key=lambda item: item["score"], reverse=True)
    summary_rows = "\n".join(
        "<tr>"
        f"<td>{escape(item['variant'])}</td>"
        f"<td class='num'>{item['pass']}</td>"
        f"<td class='num'>{item['review']}</td>"
        f"<td class='num'>{item['fail']}</td>"
        f"<td class='score'>{item['score']}</td>"
        f"<td><span class='pill {item['status']}'>{item['status']}</span></td>"
        "</tr>"
        for item in summaries
    )
    criteria_rows = "\n".join(
        "<tr>"
        f"<td>{escape(item['criterion'])}</td>"
        f"<td>{escape(item['direction'])}</td>"
        f"<td class='input'>{item['threshold']}</td>"
        f"<td>{item['weight']:.0%}</td>"
        "</tr>"
        for item in CRITERIA
    )
    matrix_rows = "\n".join(
        "<tr>"
        f"<td>{escape(record['variant'])}</td>"
        f"<td>{escape(record['criterion'])}</td>"
        f"<td class='num'>{record['value']}</td>"
        f"<td>{escape(record['unit'])}</td>"
        f"<td><span class='pill {escape(record['pass_fail'])}'>{escape(record['pass_fail'])}</span></td>"
        "</tr>"
        for record in records[:12]
    )
    PREVIEW.write_text(
        f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<style>
body {{ margin: 0; font-family: Inter, Arial, sans-serif; background: #f7f4ee; color: #243447; }}
.sheet {{ padding: 16px; }}
.top {{ display: grid; grid-template-columns: 1fr 1fr; gap: 14px; }}
.panel {{ background: white; border: 1px solid #ddd5c8; box-shadow: 0 1px 4px rgba(0,0,0,.06); }}
.tab {{ display: inline-block; background: #243447; color: white; padding: 6px 10px; font-size: 12px; font-weight: 700; letter-spacing: .04em; text-transform: uppercase; }}
h1 {{ margin: 10px 0 6px; font-size: 18px; }}
p {{ margin: 0 0 10px; color: #5d6872; font-size: 13px; }}
table {{ width: 100%; border-collapse: collapse; font-size: 12px; }}
th {{ background: #243447; color: white; text-align: left; font-weight: 700; padding: 7px 8px; }}
td {{ border: 1px solid #e6dfd4; padding: 6px 8px; vertical-align: top; }}
.num, .score {{ text-align: right; font-variant-numeric: tabular-nums; }}
.score {{ font-weight: 800; }}
.input {{ background: #fff2cc; font-weight: 700; text-align: right; }}
.pill {{ display: inline-block; min-width: 58px; text-align: center; border-radius: 999px; padding: 2px 8px; font-size: 11px; font-weight: 800; text-transform: uppercase; }}
.pass, .candidate {{ background: #d9ead3; color: #2f6336; }}
.review {{ background: #fce5cd; color: #8a4f13; }}
.fail, .revise {{ background: #f4cccc; color: #8b2722; }}
.wide {{ margin-top: 14px; }}
.note {{ margin-top: 10px; padding: 8px 10px; background: #eaf2f8; border-left: 4px solid #4b6f9f; font-size: 12px; }}
</style>
</head>
<body>
<div class="sheet">
  <span class="tab">live_threshold_workbook.xlsx</span>
  <h1>Week 3 threshold workbook preview</h1>
  <p>Static slide preview. Open the workbook for editable formulas, color rules, and chart updates.</p>
  <div class="top">
    <div class="panel">
      <table>
        <thead><tr><th>criterion</th><th>rule</th><th>threshold</th><th>weight</th></tr></thead>
        <tbody>{criteria_rows}</tbody>
      </table>
    </div>
    <div class="panel">
      <table>
        <thead><tr><th>variant</th><th>pass</th><th>review</th><th>fail</th><th>score</th><th>status</th></tr></thead>
        <tbody>{summary_rows}</tbody>
      </table>
    </div>
  </div>
  <div class="panel wide">
    <table>
      <thead><tr><th>variant</th><th>criterion</th><th>value</th><th>unit</th><th>status</th></tr></thead>
      <tbody>{matrix_rows}</tbody>
    </table>
  </div>
  <div class="note">Teaching move: change a yellow threshold or score rule, then ask which design recommendation changes and whether that threshold was evidence or judgment.</div>
</div>
</body>
</html>
""",
        encoding="utf-8",
    )


if __name__ == "__main__":
    build_workbook()
    print(f"Wrote {OUTPUT}")
    print(f"Wrote {PREVIEW}")
